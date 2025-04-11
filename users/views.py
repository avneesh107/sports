from django.shortcuts import render,redirect
from django.contrib.auth import logout

from .forms import UserProfileForm
from .forms import MatchUploadForm
import openpyxl
from django.contrib.auth.models import User
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth import login, authenticate, logout
from django.contrib import messages
from .models import Event, Team, College,UserProfile,User,Match,Result,Result_Many

from .tables import TeamTable,ParticipantTable,MatchTable

from django_tables2 import RequestConfig
from django.contrib.auth.decorators import user_passes_test




def register(request):
    if request.user.is_authenticated:
        profile = UserProfile.objects.filter(user=request.user).first()
        if not profile:
            if request.method == "POST":
                form = UserProfileForm(request.POST)
                if form.is_valid():
                    profile = form.save(commit=False)
                    profile.user = request.user
                    profile.save()
                    return redirect('dashboard')
            else:
                form = UserProfileForm()
            return render(request, 'users/register.html', {'form': form})
        else:
            return redirect('dashboard')
    else:
        return render(request, 'users/home.html')


def dashboard(request):
    if request.user.is_authenticated:
        if request.user.is_superuser:
            return redirect('admin_dashboard')
        else:
            profile = UserProfile.objects.filter(user=request.user).first()
            if not profile:
                return redirect('register')
            return render(request, 'users/dashboard.html', {'profile': profile})
       
    else:
        return render(request, 'users/home.html')


def home(request):
    if request.user.is_authenticated:
        return redirect('dashboard')  
    return render(request, 'users/home.html')  


def logout_view(request):
    logout(request)
    return redirect("/")


def event_list(request):
    events = Event.objects.all()
    teams=Team.objects.all()
    profile = UserProfile.objects.filter(user=request.user).first()
    return render(request, 'users/event_list.html', {'events': events,'teams':teams,'profile':profile})


def join_event(request, event_id):
    user = request.user.userprofile 
    event = get_object_or_404(Event, id=event_id)
    if event.gender != "X" and user.gender != event.gender:
        messages.error(request, f"You cannot join an event of a different gender.{user.gender}")
        return redirect("dashboard")
    
    team = Team.objects.filter(event=event,college=user.college).first()
    if team == None:
        team = Team.objects.create(event=event, college=user.college)
    
    if user in team.members.all():
        messages.error(request, f"You are already in the {event.sport} event.")
        return redirect("dashboard")

    if team.is_full ==0:
        messages.error(request, "This team is already full.")
        return redirect("dashboard")
    
    team.members.add(user)
    team.save()

    messages.success(request, f"You have joined the {event.sport} event.")
    return redirect("dashboard")  

def match_list(request):
    
    matches = Match.objects.all()

    return render(request, 'users/match_list.html', {'matches': matches})

@user_passes_test(lambda u: u.is_superuser)
def team_list(request):
    teams = Team.objects.select_related('event', 'college')  
    table = TeamTable(teams)
    RequestConfig(request).configure(table)
    return render(request, 'users/team_list.html', {'table': table})


def team_members(request, pk):
    team = Team.objects.get(id=pk)
    members = team.members.all()  
    return render(request, 'users/team_members.html', {'team': team, 'members': members})

def admin_dashboard(request):
    table = ParticipantTable(UserProfile.objects.select_related('college'))
    RequestConfig(request).configure(table)
    user=UserProfile.objects.all()
    college=College.objects.all()
    return render(request, "users/admin_dashboard.html", {"table": table,'user':user,'college':college})


def college_participants_view(request,college_id):
    college = get_object_or_404(College, id=college_id)
    participants = UserProfile.objects.filter(college=college)
    return render(request, "users/college_participants.html", {
        "college": college,
        "participants": participants
    })

def Matches(request):
    table = MatchTable(Match.objects.all())
    RequestConfig(request).configure(table)
    
    return render(request, "users/Match.html", {"table": table})

def cont_team(request,pk):
    match = get_object_or_404(Match,id=pk) 
    teams=match.contestants.all()
    table = TeamTable(teams)
    RequestConfig(request).configure(table)
    return render(request, 'users/team_list.html', {'table': table})

@user_passes_test(lambda u: u.is_superuser)
def result(request):
    team_result= Result.objects.all()
    many_result=Result_Many.objects.all()
    return render(request,'users/result.html',{'team_results':team_result,'many_results':many_result})


def User_Matches(request):
    profile = UserProfile.objects.filter(user=request.user).first()
    college=profile.college
    matches = Match.objects.filter(contestants__college=college).distinct()
    
    table = MatchTable(matches)
    RequestConfig(request).configure(table)
    
    return render(request, "users/Match.html", {"table": table})


def user_results_view(request):
    profile = UserProfile.objects.filter(user=request.user).first()
    college=profile.college
    team_result = Result.objects.filter(team__college=college)
    many_result = Result_Many.objects.filter(team__college=college)

    return render(request,'users/result.html',{'team_results':team_result,'many_results':many_result})

@user_passes_test(lambda u: u.is_superuser)
def bulk_match_upload_view(request):
    if request.method == "POST":
        form = MatchUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = form.cleaned_data['file']
            wb = openpyxl.load_workbook(file)
            sheet = wb.active

            for row in sheet.iter_rows(min_row=2, values_only=True):  
                try:
                    event_name, date, start_time, end_time, team_ids_str = row
                    event = Event.objects.get(name=event_name)
                    match = Match.objects.create(
                        event=event,
                        date=date,
                        start=start_time,
                        end=end_time,
                    )
                    team_ids = [int(t.strip()) for t in str(team_ids_str).split(",")]
                    teams = Team.objects.filter(id__in=team_ids)
                    match.contestants.set(teams)
                    match.save()
                except Exception as e:
                    messages.error(request, f"Error in row: {row} — {e}")
                    continue

            messages.success(request, "Matches uploaded successfully!")
            return redirect("admin_dashboard")

    else:
        form = MatchUploadForm()

    return render(request, "users/upload_excel.html", {"form": form})